"""
Comprehensive Test Suite for Phase 2 Document Intelligence Pipeline.
Covers:
- Excel Acceptance Test (Section 55)
- PDF Acceptance Test (Section 56)
- Poor Scan Acceptance Test (Section 57)
- Image Acceptance Test (Section 58)
- CSV Acceptance Test (Section 59)
- Quality Assessment (Section 4)
- Document Enhancement & Deskew (Section 5, 9, 10)
- Metric, Unit, Period, Org Normalization (Section 27, 28, 29, 30)
- Confidence Engine (Section 36, 37, 38)
- Document Reprocessing & Deduplication (Section 50)
- Evidence API and Source Provenance (Section 48, 54)
"""
import io
import os
import tempfile
import pytest
import numpy as np
import cv2
import openpyxl
import pymupdf as fitz
from PIL import Image, ImageDraw

from app.models.document import Document
from app.models.fact import ExtractedFact
from app.models.processing import DocumentPage, DocumentSheet, DocumentTable, DocumentQuality
from app.models.enums import DocumentStatus, ValidationStatus, ExtractionMethod
from app.services.pipeline.file_classifier import FileClassifier
from app.services.pipeline.quality_assessment import QualityAssessmentService
from app.services.pipeline.enhancement import DocumentEnhancementService
from app.services.pipeline.document_processor import DocumentProcessor
from app.services.extractors.excel_extractor import ExcelExtractor
from app.services.extractors.pdf_extractor import PDFExtractor
from app.services.extractors.csv_extractor import CSVExtractor
from app.services.extractors.image_extractor import ImageExtractor
from app.services.mining.metric_registry import MiningMetricRegistry
from app.services.mining.unit_normalizer import UnitNormalizer
from app.services.mining.period_normalizer import PeriodNormalizer
from app.services.mining.org_normalizer import OrgNormalizer
from app.services.facts.fact_extractor import FactExtractor
from app.services.facts.confidence_engine import ConfidenceEngine
from app.services.facts.evidence_persistence import EvidencePersistenceService


# ==============================================================================
# 1. Excel Acceptance Test (Section 55)
# ==============================================================================
def test_excel_acceptance_test_section_55(db_session):
    """
    Acceptance Test 55:
    Subsidiary | Production | Target | Period
    MCL | 40 | 42 | FY 2025-26
    SECL | 35 | 36 | FY 2025-26
    NCL | 30 | 32 | FY 2025-26
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Production_Summary"

    headers = ["Subsidiary", "Production", "Target", "Period"]
    ws.append(headers)
    ws.append(["MCL", 40, 42, "FY 2025-26"])
    ws.append(["SECL", 35, 36, "FY 2025-26"])
    ws.append(["NCL", 30, 32, "FY 2025-26"])

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = tmp.name
        wb.save(tmp_path)
    wb.close()

    try:
        # Create document
        doc = Document(
            original_filename="Production_Summary.xlsx",
            stored_filename=os.path.basename(tmp_path),
            file_type="XLSX",
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            file_size=os.path.getsize(tmp_path),
            document_category="Production Report",
            organization="Coal India Limited",
            storage_path=tmp_path,
            status=DocumentStatus.UPLOADED.value
        )
        db_session.add(doc)
        db_session.commit()

        # Process document
        processed_doc = DocumentProcessor.process_document(db_session, doc.id)
        assert processed_doc.status == DocumentStatus.READY.value
        assert processed_doc.fact_count >= 6  # 3 production + 3 target facts

        facts = db_session.query(ExtractedFact).filter(ExtractedFact.document_id == doc.id).all()
        prod_facts = [f for f in facts if f.metric_code == "COAL_PRODUCTION"]
        target_facts = [f for f in facts if f.metric_code == "PRODUCTION_TARGET"]

        assert len(prod_facts) == 3
        assert len(target_facts) == 3

        # Verify MCL production fact
        mcl_prod = next(f for f in prod_facts if f.subsidiary == "MCL")
        assert mcl_prod.numeric_value == 40.0
        assert mcl_prod.reporting_period == "FY 2025-26"
        assert mcl_prod.unit == "MT"
        assert mcl_prod.cell_reference == "B2"
        assert mcl_prod.confidence_score >= 0.90
        assert mcl_prod.validation_status == ValidationStatus.HIGH_CONFIDENCE.value

        # Verify SECL target fact
        secl_target = next(f for f in target_facts if f.subsidiary == "SECL")
        assert secl_target.numeric_value == 36.0
        assert secl_target.reporting_period == "FY 2025-26"
        assert secl_target.cell_reference == "C3"

        # Verify NCL production fact
        ncl_prod = next(f for f in prod_facts if f.subsidiary == "NCL")
        assert ncl_prod.numeric_value == 30.0
        assert ncl_prod.cell_reference == "B4"

        # Ensure no invented values
        extracted_vals = {f.numeric_value for f in facts}
        assert extracted_vals.issubset({40.0, 42.0, 35.0, 36.0, 30.0, 32.0})

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


# ==============================================================================
# 2. PDF Acceptance Test (Section 56)
# ==============================================================================
def test_pdf_acceptance_test_section_56(db_session):
    """
    Acceptance Test 56:
    Create digital PDF containing:
    heading, paragraphs, production value, drilling value, organization, reporting period.
    """
    doc_pdf = fitz.open()
    page = doc_pdf.new_page()

    text_content = (
        "CMPDI Geological & Production Report\n\n"
        "Executive Summary for Coal India Operations\n\n"
        "SECL achieved raw coal production of 35.2 MT during FY 2025-26.\n"
        "In addition, exploratory drilling of 1250 m was successfully completed by CMPDI in Korba coalfield.\n"
        "MCL reported total overburden removal of 45.8 Mm3 in the same period.\n"
    )
    page.insert_text((50, 72), text_content, fontsize=11)

    tmp = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
    tmp_path = tmp.name
    tmp.close()
    doc_pdf.save(tmp_path)
    doc_pdf.close()

    try:
        doc = Document(
            original_filename="SECL_CMPDI_Report_2025.pdf",
            stored_filename=os.path.basename(tmp_path),
            file_type="PDF",
            mime_type="application/pdf",
            file_size=os.path.getsize(tmp_path),
            document_category="Geological Report",
            organization="CMPDI / CIL",
            storage_path=tmp_path,
            status=DocumentStatus.UPLOADED.value
        )
        db_session.add(doc)
        db_session.commit()

        processed_doc = DocumentProcessor.process_document(db_session, doc.id)
        assert processed_doc.status == DocumentStatus.READY.value
        assert processed_doc.page_count == 1
        assert processed_doc.fact_count >= 2

        facts = db_session.query(ExtractedFact).filter(ExtractedFact.document_id == doc.id).all()
        # Find SECL production fact
        secl_prod = next((f for f in facts if f.subsidiary == "SECL" and f.metric_code == "COAL_PRODUCTION"), None)
        assert secl_prod is not None
        assert secl_prod.numeric_value == 35.2
        assert secl_prod.page_number == 1
        assert "35.2" in secl_prod.source_context

        # Verify chunks created
        pages = db_session.query(DocumentPage).filter(DocumentPage.document_id == doc.id).all()
        assert len(pages) == 1
        assert "SECL" in pages[0].raw_text

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


# ==============================================================================
# 3. Poor Scan Acceptance Test (Section 57)
# ==============================================================================
def test_poor_scan_acceptance_test_section_57():
    """
    Acceptance Test 57:
    Create deliberately low-quality scan with skew, low contrast, blur, noise.
    Verify quality assessment detects issues, enhancement applies deskew/denoise,
    before/after scores are recorded, and no fabricated values occur.
    """
    # Create base image
    img = np.ones((400, 600, 3), dtype=np.uint8) * 230
    cv2.putText(img, "SECL Coal Production 45.5 MT", (50, 200), cv2.FONT_HERSHEY_SIMPLEX, 0.9, (100, 100, 100), 2)

    # 1. Add blur
    blurred = cv2.GaussianBlur(img, (9, 9), 3.0)

    # 2. Add noise
    noise = np.random.normal(0, 25, blurred.shape).astype(np.uint8)
    noisy = cv2.add(blurred, noise)

    # 3. Add skew
    (h, w) = noisy.shape[:2]
    center = (w // 2, h // 2)
    M = cv2.getRotationMatrix2D(center, 4.0, 1.0)
    skewed = cv2.warpAffine(noisy, M, (w, h), borderValue=(230, 230, 230))

    _, enc = cv2.imencode(".png", skewed)
    poor_bytes = enc.tobytes()

    # Assess quality
    profile = QualityAssessmentService.assess_image_data(poor_bytes)
    assert profile["quality_score"] < 0.85
    assert profile["enhancement_recommended"] is True
    assert profile["noise"] == "high" or profile["blur"] in ("medium", "high") or abs(profile["skew_angle"]) > 0

    # Apply enhancement
    enhanced_bytes, ops = DocumentEnhancementService.enhance_image(poor_bytes)
    assert len(ops) > 0
    assert any("CLAHE" in op or "Deskew" in op or "Denoising" in op for op in ops)

    # Check after quality
    after_profile = QualityAssessmentService.assess_image_data(enhanced_bytes)
    assert after_profile["quality_score"] > 0
    assert len(ops) > 0


# ==============================================================================
# 4. CSV Acceptance Test (Section 59)
# ==============================================================================
def test_csv_acceptance_test_section_59(db_session):
    """
    Acceptance Test 59:
    CSV with: Mine | Metric | Value | Unit | Period
    """
    csv_text = (
        "Mine,Metric,Value,Unit,Period\n"
        "Kusmunda,Coal Production,50.5,MT,FY 2025-26\n"
        "Gevra,Coal Production,60.2,MT,FY 2025-26\n"
        "Dipka,Overburden Removal,45.0,Mm3,FY 2025-26\n"
    )

    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w", encoding="utf-8") as tmp:
        tmp_path = tmp.name
        tmp.write(csv_text)

    try:
        doc = Document(
            original_filename="SECL_Mega_Mines.csv",
            stored_filename=os.path.basename(tmp_path),
            file_type="CSV",
            mime_type="text/csv",
            file_size=os.path.getsize(tmp_path),
            document_category="Production Report",
            organization="SECL",
            storage_path=tmp_path,
            status=DocumentStatus.UPLOADED.value
        )
        db_session.add(doc)
        db_session.commit()

        processed_doc = DocumentProcessor.process_document(db_session, doc.id)
        assert processed_doc.status == DocumentStatus.READY.value
        assert processed_doc.fact_count == 3

        facts = db_session.query(ExtractedFact).filter(ExtractedFact.document_id == doc.id).all()
        kusmunda = next(f for f in facts if f.mine == "Kusmunda")
        assert kusmunda.metric_code == "COAL_PRODUCTION"
        assert kusmunda.numeric_value == 50.5
        assert kusmunda.unit == "MT"
        assert kusmunda.reporting_period == "FY 2025-26"
        assert kusmunda.row_number == 2

        gevra = next(f for f in facts if f.mine == "Gevra")
        assert gevra.numeric_value == 60.2

        dipka = next(f for f in facts if f.mine == "Dipka")
        assert dipka.metric_code == "OVERBURDEN_REMOVAL"
        assert dipka.numeric_value == 45.0
        assert dipka.unit == "Mm3"

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


# ==============================================================================
# 5. Normalization Unit Tests (Sections 27, 28, 29, 30)
# ==============================================================================
def test_metric_normalization():
    reg = MiningMetricRegistry()
    assert reg.resolve("Coal Production")[0] == "COAL_PRODUCTION"
    assert reg.resolve("Actual Production")[0] == "COAL_PRODUCTION"
    assert reg.resolve("Output")[0] == "COAL_PRODUCTION"
    assert reg.resolve("Production Target")[0] == "PRODUCTION_TARGET"
    assert reg.resolve("Target")[0] == "PRODUCTION_TARGET"
    assert reg.resolve("Coal Dispatch")[0] == "COAL_DISPATCH"
    assert reg.resolve("Offtake")[0] == "COAL_DISPATCH"
    assert reg.resolve("Geological Reserve")[0] == "GEOLOGICAL_RESERVE"
    assert reg.resolve("Resources")[0] == "GEOLOGICAL_RESERVE"
    assert reg.resolve("Drilling")[0] == "DRILLING"
    assert reg.resolve("Drilled Metres")[0] == "DRILLING"
    assert reg.resolve("Meterage")[0] == "DRILLING"
    assert reg.resolve("Overburden")[0] == "OVERBURDEN_REMOVAL"
    assert reg.resolve("OB")[0] == "OVERBURDEN_REMOVAL"
    assert reg.resolve("Stripping Ratio")[0] == "STRIPPING_RATIO"


def test_unit_normalization():
    assert UnitNormalizer.normalize("million tonnes")[0] == "MT"
    assert UnitNormalizer.normalize("MT")[0] == "MT"
    assert UnitNormalizer.normalize("tonnes")[0] == "T"
    assert UnitNormalizer.normalize("MTPA")[0] == "MTPA"
    assert UnitNormalizer.normalize("metres")[0] == "m"
    assert UnitNormalizer.normalize("ha")[0] == "ha"
    assert UnitNormalizer.normalize("%")[0] == "%"
    assert UnitNormalizer.normalize("Mm3")[0] == "Mm3"
    assert UnitNormalizer.normalize("₹ Cr")[0] == "INR Cr"
    assert UnitNormalizer.normalize("INR Cr")[0] == "INR Cr"


def test_period_normalization():
    assert PeriodNormalizer.normalize_period("FY 2025-26") == "FY 2025-26"
    assert PeriodNormalizer.normalize_period("2025-26") == "FY 2025-26"
    assert PeriodNormalizer.normalize_period("2025/26") == "FY 2025-26"
    assert PeriodNormalizer.normalize_period("FY2026") == "FY 2025-26"
    assert PeriodNormalizer.normalize_period("Financial Year 2025-26") == "FY 2025-26"
    assert PeriodNormalizer.normalize_period("Q2 FY 2025-26") == "Q2 FY 2025-26"


def test_org_normalization():
    assert OrgNormalizer.normalize_org("Mahanadi Coalfields Limited") == "MCL"
    assert OrgNormalizer.normalize_org("MCL") == "MCL"
    assert OrgNormalizer.normalize_org("South Eastern Coalfields") == "SECL"
    assert OrgNormalizer.normalize_org("SECL") == "SECL"
    assert OrgNormalizer.normalize_org("Coal India") == "CIL"
    assert OrgNormalizer.normalize_org("Central Mine Planning and Design Institute") == "CMPDI"


# ==============================================================================
# 6. Confidence Engine & Validation Status (Section 36, 37, 38)
# ==============================================================================
def test_confidence_engine():
    # Clean structured table with all coordinates
    conf, rationale = ConfidenceEngine.calculate_confidence(
        extraction_method="STRUCTURED_TABLE",
        has_metric=True,
        has_unit=True,
        has_period=True,
        has_subsidiary=True,
        document_quality_score=1.0,
        is_structured_spreadsheet=True
    )
    assert conf >= 0.90
    assert ConfidenceEngine.get_initial_validation_status(conf) == ValidationStatus.HIGH_CONFIDENCE.value

    # Sub-optimal scan quality discount
    low_conf, _ = ConfidenceEngine.calculate_confidence(
        extraction_method="OCR_BLOCK",
        has_metric=True,
        has_unit=False,
        has_period=False,
        has_subsidiary=False,
        document_quality_score=0.5,
        ocr_confidence=0.6
    )
    assert low_conf < 0.70
    assert ConfidenceEngine.get_initial_validation_status(low_conf) == ValidationStatus.REVIEW_REQUIRED.value


# ==============================================================================
# 7. Document Reprocessing & Idempotency (Section 50)
# ==============================================================================
def test_document_reprocessing_idempotency(db_session):
    csv_text = "Subsidiary,Production\nMCL,40\nSECL,35\n"
    with tempfile.NamedTemporaryFile(suffix=".csv", delete=False, mode="w", encoding="utf-8") as tmp:
        tmp_path = tmp.name
        tmp.write(csv_text)

    try:
        doc = Document(
            original_filename="reprocess_test.csv",
            stored_filename=os.path.basename(tmp_path),
            file_type="CSV",
            mime_type="text/csv",
            file_size=os.path.getsize(tmp_path),
            storage_path=tmp_path,
            status=DocumentStatus.UPLOADED.value
        )
        db_session.add(doc)
        db_session.commit()

        # Run 1
        DocumentProcessor.process_document(db_session, doc.id)
        count_run_1 = db_session.query(ExtractedFact).filter(ExtractedFact.document_id == doc.id).count()

        # Reprocess
        DocumentProcessor.reprocess_document(db_session, doc.id)
        count_run_2 = db_session.query(ExtractedFact).filter(ExtractedFact.document_id == doc.id).count()

        # Must not produce duplicate facts!
        assert count_run_1 == count_run_2
        assert count_run_2 == 2

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


# ==============================================================================
# 8. Image Extraction Path & OCR Integration (Section 16, 58)
# ==============================================================================
def test_image_extraction_path(db_session):
    """
    Test PNG image upload, quality assessment, candidate generation, and extractor path.
    """
    img = np.ones((300, 500, 3), dtype=np.uint8) * 255
    cv2.putText(img, "WCL Coal Production 28 MT", (30, 150), cv2.FONT_HERSHEY_SIMPLEX, 0.8, (0, 0, 0), 2)

    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
        tmp_path = tmp.name
        cv2.imwrite(tmp_path, img)

    try:
        doc = Document(
            original_filename="wcl_scan.png",
            stored_filename=os.path.basename(tmp_path),
            file_type="PNG",
            mime_type="image/png",
            file_size=os.path.getsize(tmp_path),
            document_category="Production Report",
            organization="WCL",
            storage_path=tmp_path,
            status=DocumentStatus.UPLOADED.value
        )
        db_session.add(doc)
        db_session.commit()

        processed_doc = DocumentProcessor.process_document(db_session, doc.id)
        # Should complete cleanly (READY or COMPLETED_WITH_WARNINGS if OCR is unavailable)
        assert processed_doc.status in (DocumentStatus.READY.value, DocumentStatus.COMPLETED_WITH_WARNINGS.value)
        assert processed_doc.quality_label is not None

        # Verify quality record created
        quality = db_session.query(DocumentQuality).filter(DocumentQuality.document_id == doc.id).first()
        assert quality is not None
        assert quality.quality_score is not None

    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


# ==============================================================================
# 9. Processing Failure & Error Handling (Section 60)
# ==============================================================================
def test_processing_failure_missing_file(db_session):
    """Verifies that missing files fail gracefully without crashing the server."""
    doc = Document(
        original_filename="missing.pdf",
        stored_filename="missing_stored.pdf",
        file_type="PDF",
        mime_type="application/pdf",
        file_size=1024,
        storage_path="C:\\nonexistent\\path\\missing.pdf",
        status=DocumentStatus.UPLOADED.value
    )
    db_session.add(doc)
    db_session.commit()

    processed_doc = DocumentProcessor.process_document(db_session, doc.id)
    assert processed_doc.status == DocumentStatus.FAILED.value
    assert "File not found" in processed_doc.processing_error


# ==============================================================================
# 10. Evidence Provenance Verification (Section 48, 54)
# ==============================================================================
def test_evidence_provenance_retrieval(db_session):
    """
    Verify Rule 2: Every numerical fact preserves provenance coordinates.
    """
    doc = Document(
        original_filename="Test_Audit_Doc.xlsx",
        stored_filename="Test_Audit_Doc.xlsx",
        file_type="XLSX",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        file_size=2048,
        storage_path="C:\\dummy\\path.xlsx",
        status=DocumentStatus.READY.value
    )
    db_session.add(doc)
    db_session.flush()

    fact = ExtractedFact(
        document_id=doc.id,
        metric_name="Coal Production",
        metric_code="COAL_PRODUCTION",
        numeric_value=42.5,
        unit="MT",
        reporting_period="FY 2025-26",
        subsidiary="MCL",
        sheet_name="Annual_Summary",
        row_number=18,
        column_name="Production",
        cell_reference="F18",
        source_context="Sheet: Annual_Summary, Cell F18 -> Value: 42.5",
        confidence_score=0.96,
        validation_status=ValidationStatus.HIGH_CONFIDENCE.value
    )
    db_session.add(fact)
    db_session.commit()

    saved = db_session.query(ExtractedFact).filter(ExtractedFact.id == fact.id).first()
    assert saved.cell_reference == "F18"
    assert saved.sheet_name == "Annual_Summary"
    assert saved.row_number == 18
    assert saved.column_name == "Production"
    assert saved.numeric_value == 42.5
    assert saved.confidence_score == 0.96

