"""
DOCX Extractor — Word document paragraph and table extraction using python-docx.
Extracts structured paragraphs, headings, and tables with full row and column provenance.
"""
import os
from typing import Dict, Any, List
from app.core.logging import logger

try:
    import docx
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False


class DOCXExtractor:
    @staticmethod
    def extract(file_path: str) -> Dict[str, Any]:
        if not _DOCX_AVAILABLE:
            raise RuntimeError("python-docx is not installed. Please install it with 'pip install python-docx'.")

        doc = docx.Document(file_path)
        
        paragraphs_text: List[str] = []
        full_text_parts: List[str] = []
        
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                # Check for heading style
                style_name = p.style.name if p.style else ""
                if "Heading" in style_name:
                    paragraphs_text.append(f"## {text}")
                    full_text_parts.append(f"## {text}")
                else:
                    paragraphs_text.append(text)
                    full_text_parts.append(text)

        # Extract Tables
        tables_data: List[Dict[str, Any]] = []
        for t_idx, table in enumerate(doc.tables):
            if not table.rows:
                continue

            # First row as headers
            headers = [cell.text.strip() or f"Col_{c_i}" for c_i, cell in enumerate(table.rows[0].cells)]
            # Deduplicate headers if empty/repeated
            seen_headers = {}
            clean_headers = []
            for h in headers:
                count = seen_headers.get(h, 0)
                if count > 0:
                    clean_headers.append(f"{h}_{count}")
                else:
                    clean_headers.append(h)
                seen_headers[h] = count + 1

            rows_data: List[Dict[str, str]] = []
            for r_idx, row in enumerate(table.rows[1:]):
                row_dict = {}
                for c_idx, cell in enumerate(row.cells):
                    col_name = clean_headers[c_idx] if c_idx < len(clean_headers) else f"Col_{c_idx}"
                    row_dict[col_name] = cell.text.strip()
                rows_data.append(row_dict)

            # Reconstruct table markdown text for full_text
            table_md_lines = [f"\n### Table {t_idx + 1}"]
            table_md_lines.append("| " + " | ".join(clean_headers) + " |")
            table_md_lines.append("| " + " | ".join(["---"] * len(clean_headers)) + " |")
            for r in rows_data:
                table_md_lines.append("| " + " | ".join([r.get(h, "") for h in clean_headers]) + " |")
            table_md = "\n".join(table_md_lines)
            full_text_parts.append(table_md)

            tables_data.append({
                "table_index": t_idx,
                "page_number": 1,  # DOCX doesn't have native physical pages without rendering engine
                "sheet_name": f"Table_{t_idx + 1}",
                "title": f"Table {t_idx + 1}",
                "source_range": f"R1C1:R{len(table.rows)}C{len(clean_headers)}",
                "row_count": len(rows_data),
                "col_count": len(clean_headers),
                "headers": clean_headers,
                "data": rows_data,
                "confidence": 1.0,
                "structure_status": "CLEAN"
            })

        full_content = "\n\n".join(full_text_parts)
        # Logical page chunks (~3000 chars)
        page_chunks = [full_content[i:i+3000] for i in range(0, max(1, len(full_content)), 3000)]
        pages_data = []

        for p_idx, p_text in enumerate(page_chunks):
            pages_data.append({
                "page_number": p_idx + 1,
                "text": p_text,
                "char_count": len(p_text),
                "word_count": len(p_text.split()),
                "has_native_text": True,
                "ocr_applied": False,
                "ocr_confidence": None,
                "image_count": 0,
                "extraction_method": "DOCX_NATIVE"
            })

        logger.info(f"Extracted DOCX '{os.path.basename(file_path)}': {len(pages_data)} logical pages, {len(tables_data)} tables, {len(full_content)} chars.")
        return {
            "page_count": len(pages_data),
            "total_chars": len(full_content),
            "is_scanned": False,
            "pages": pages_data,
            "tables": tables_data,
            "text": full_content
        }
